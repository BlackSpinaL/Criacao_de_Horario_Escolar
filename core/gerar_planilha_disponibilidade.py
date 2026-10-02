"""Gera planilha de disponibilidade em branco (formato CTPM/Lavras)."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DIAS = ["SEGUNDA", "TERÇA", "QUARTA", "QUINTA", "SEXTA"]

COR_HEADER = PatternFill("solid", fgColor="1F4E78")
COR_VESP   = PatternFill("solid", fgColor="FCE4D6")
COR_MAT    = PatternFill("solid", fgColor="DDEBF7")
COR_AULA   = PatternFill("solid", fgColor="FFF9E6")
COR_OBS    = PatternFill("solid", fgColor="F2F2F2")

FONTE_HEADER = Font(color="FFFFFF", bold=True, size=10)
FONTE_TIT    = Font(bold=True, size=11, color="1F4E78")
FONTE_BOLD   = Font(bold=True, size=10)
FONTE_OBS    = Font(italic=True, size=9, color="7F8C8D")

BORDA_FINA = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
CENTRO = Alignment(horizontal="center", vertical="center")
ESQ = Alignment(horizontal="left", vertical="center")


def _desenhar_bloco(ws, linha_inicio, titulo, n_aulas, cor_titulo):
    row = linha_inicio

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    c = ws.cell(row=row, column=1, value=titulo)
    c.font = FONTE_TIT
    c.fill = cor_titulo
    c.alignment = CENTRO
    c.border = BORDA_FINA
    row += 1

    c = ws.cell(row=row, column=1, value="HORÁRIO")
    c.font = FONTE_HEADER
    c.fill = COR_HEADER
    c.alignment = CENTRO
    c.border = BORDA_FINA
    for j, dia in enumerate(DIAS, start=2):
        c = ws.cell(row=row, column=j, value=dia)
        c.font = FONTE_HEADER
        c.fill = COR_HEADER
        c.alignment = CENTRO
        c.border = BORDA_FINA
    row += 1

    for i in range(1, n_aulas + 1):
        c = ws.cell(row=row, column=1, value=f"{i}º horário")
        c.font = FONTE_BOLD
        c.fill = cor_titulo
        c.alignment = CENTRO
        c.border = BORDA_FINA
        for j in range(2, 7):
            c = ws.cell(row=row, column=j)
            c.fill = COR_AULA
            c.alignment = CENTRO
            c.border = BORDA_FINA
        row += 1

    return row + 1


def gerar_planilha_bytes(nome_professor, disciplinas=None):
    """Gera a planilha em memória e devolve os bytes do .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Disponibilidade"

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"Professor:  {nome_professor}"
    c.font = Font(bold=True, size=12, color="1F4E78")
    c.alignment = ESQ

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    disc_str = ", ".join(disciplinas) if disciplinas else "(a preencher)"
    c.value = f"Disciplinas:  {disc_str}"
    c.font = FONTE_BOLD
    c.alignment = ESQ

    prox = _desenhar_bloco(
        ws, linha_inicio=3,
        titulo="VESPERTINO – 6 AULAS DE 45 MINUTOS",
        n_aulas=6,
        cor_titulo=COR_VESP,
    )

    prox = _desenhar_bloco(
        ws, linha_inicio=prox,
        titulo="MATUTINO – 7 AULAS DE 45 MINUTOS",
        n_aulas=7,
        cor_titulo=COR_MAT,
    )

    ws.merge_cells(start_row=prox, start_column=1, end_row=prox, end_column=6)
    c = ws.cell(row=prox, column=1,
                value="OBSERVAÇÃO IMPORTANTE: O(a) professor(a) deverá indicar "
                      "25% a mais de disponibilidade de horários, a fim de "
                      "facilitar a elaboração.")
    c.font = FONTE_OBS
    c.fill = COR_OBS
    c.alignment = ESQ
    prox += 1

    ws.merge_cells(start_row=prox, start_column=1, end_row=prox, end_column=6)
    c = ws.cell(row=prox, column=1,
                value="Orientação: Pinte todos os horários em que possui "
                      "disponibilidade, considerando também os 25% adicionais.")
    c.font = FONTE_OBS
    c.fill = COR_OBS
    c.alignment = ESQ
    prox += 2

    ws.merge_cells(start_row=prox, start_column=1, end_row=prox + 3, end_column=6)
    c = ws.cell(row=prox, column=1, value="Observações:")
    c.font = FONTE_BOLD
    c.alignment = Alignment(horizontal="left", vertical="top")

    ws.column_dimensions["A"].width = 18
    for j in range(2, 7):
        ws.column_dimensions[get_column_letter(j)].width = 16

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
