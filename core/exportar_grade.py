"""Exporta a grade gerada em PDF e Excel."""
from io import BytesIO
from datetime import datetime
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor, white
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak,
)
from sqlalchemy import text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


AZUL = HexColor("#1F4E78")
CINZA_CLARO = HexColor("#F0F2F6")
BORDA_CINZA = HexColor("#808080")

DIAS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]

XL_AZUL = PatternFill("solid", fgColor="1F4E78")
XL_AZUL_CLARO = PatternFill("solid", fgColor="D6EAF8")
XL_FONTE_HEADER = Font(color="FFFFFF", bold=True, size=11)
XL_FONTE_BOLD = Font(bold=True, size=10)
XL_FONTE_NORMAL = Font(size=10)
XL_BORDA = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)
XL_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _buscar_grade_turma(engine, ano_id, turma_id):
    with engine.connect() as conn:
        turma = conn.execute(text("""
            SELECT t.id, t.codigo, t.nome, t.serie, t.turno, t.segmento,
                   t.grade_horaria_id
            FROM turmas t WHERE t.id = :t
        """), {"t": turma_id}).first()
        if not turma:
            return None

        slots = conn.execute(text("""
            SELECT ordem, nome, inicio, fim
            FROM slots_horario
            WHERE grade_id = :g AND tipo = 'aula'
            ORDER BY ordem
        """), {"g": turma.grade_horaria_id}).fetchall()

        aloc = conn.execute(text("""
            SELECT h.dia, h.ordem,
                   c.nome AS componente, p.nome AS professor
            FROM grade_gerada g
            JOIN atividades a ON a.id = g.atividade_id
            JOIN horarios h ON h.id = g.horario_id
            JOIN componentes c ON c.id = a.componente_id
            JOIN professores p ON p.id = a.professor_id
            WHERE a.turma_id = :t AND g.ano_letivo_id = :a
        """), {"t": turma_id, "a": ano_id}).fetchall()

    mapa = {}
    for a in aloc:
        mapa[(a.dia, a.ordem)] = (a.componente, a.professor)

    return {"turma": turma, "slots": slots, "mapa": mapa}


def _todas_turmas(engine, ano_id):
    with engine.connect() as conn:
        return [r.id for r in conn.execute(text("""
            SELECT id FROM turmas WHERE ano_letivo_id = :a ORDER BY codigo
        """), {"a": ano_id}).fetchall()]


def _cabecalho_rodape(canvas, doc, titulo_extra=""):
    canvas.saveState()
    largura, altura = landscape(A4)
    canvas.setFillColor(AZUL)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(1.5 * cm, altura - 1 * cm,
                      "CTPM/Lavras — Sistema de Grade Horária")
    if titulo_extra:
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(largura - 1.5 * cm, altura - 1 * cm, titulo_extra)
    canvas.setStrokeColor(AZUL)
    canvas.setLineWidth(0.5)
    canvas.line(1.5 * cm, altura - 1.2 * cm, largura - 1.5 * cm, altura - 1.2 * cm)

    canvas.setFillColor(HexColor("#7F8C8D"))
    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(largura / 2, 0.7 * cm, f"Página {doc.page}")
    canvas.drawRightString(largura - 1.5 * cm, 0.7 * cm,
                          datetime.now().strftime("%d/%m/%Y"))
    canvas.restoreState()


def gerar_pdf_grade(engine, ano_id, turma_ids=None):
    """Gera PDF com uma página por turma (paisagem A4)."""
    if turma_ids is None:
        turma_ids = _todas_turmas(engine, ano_id)

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.6 * cm, bottomMargin=1.4 * cm,
        title="Grade Horária CTPM/Lavras",
    )

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "titulo_turma", parent=estilos["Title"],
        fontSize=15, textColor=AZUL, alignment=TA_CENTER, spaceAfter=4,
    )
    estilo_sub = ParagraphStyle(
        "sub", parent=estilos["Normal"],
        alignment=TA_CENTER, fontSize=9,
        textColor=HexColor("#7F8C8D"), spaceAfter=10,
    )
    estilo_celula = ParagraphStyle(
        "celula", parent=estilos["Normal"],
        fontSize=7.5, alignment=TA_CENTER, leading=9,
    )
    estilo_dia = ParagraphStyle(
        "dia", parent=estilos["Normal"],
        fontSize=9, alignment=TA_CENTER, textColor=white,
        fontName="Helvetica-Bold",
    )
    estilo_horario = ParagraphStyle(
        "horario", parent=estilos["Normal"],
        fontSize=7.5, alignment=TA_CENTER, leading=9,
        fontName="Helvetica-Bold", textColor=AZUL,
    )

    story = []
    primeira = True

    for tid in turma_ids:
        dados = _buscar_grade_turma(engine, ano_id, tid)
        if not dados:
            continue

        t = dados["turma"]
        slots = dados["slots"]
        mapa = dados["mapa"]

        if not primeira:
            story.append(PageBreak())
        primeira = False

        story.append(Paragraph(f"Grade Horária — {t.nome}", estilo_titulo))
        story.append(Paragraph(
            f"Série: {t.serie} &nbsp;|&nbsp; Turno: {t.turno} "
            f"&nbsp;|&nbsp; Código: {t.codigo}",
            estilo_sub,
        ))

        if not slots:
            story.append(Paragraph(
                "Nenhuma aula atribuída a esta turma.",
                ParagraphStyle("vazio", parent=estilos["Normal"],
                              alignment=TA_CENTER, fontSize=11),
            ))
            continue

        cabecalho = [Paragraph("Horário", estilo_dia)]
        for dia in DIAS:
            cabecalho.append(Paragraph(dia.upper(), estilo_dia))

        tabela_dados = [cabecalho]

        for slot in slots:
            linha = [Paragraph(
                f"{slot.nome}<br/>{slot.inicio}-{slot.fim}",
                estilo_horario,
            )]
            for dia in DIAS:
                celula = mapa.get((dia, slot.ordem))
                if celula:
                    comp, prof = celula
                    texto = (f"<b>{comp}</b><br/>"
                             f"<font size='6.5' color='#555555'>({prof})</font>")
                    linha.append(Paragraph(texto, estilo_celula))
                else:
                    linha.append(Paragraph("—", estilo_celula))
            tabela_dados.append(linha)

        largura_util = landscape(A4)[0] - 3 * cm
        largura_horario = 2.6 * cm
        largura_dia = (largura_util - largura_horario) / 5
        col_widths = [largura_horario] + [largura_dia] * 5

        tbl = Table(tabela_dados, colWidths=col_widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), AZUL),
            ("TEXTCOLOR", (0, 0), (-1, 0), white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, BORDA_CINZA),
            ("BACKGROUND", (0, 1), (0, -1), CINZA_CLARO),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (1, 1), (-1, -1),
             [white, HexColor("#FAFBFC")]),
        ]))
        story.append(tbl)

    if not story:
        story.append(Paragraph(
            "Nenhuma turma encontrada.",
            ParagraphStyle("vazio", parent=estilos["Normal"], fontSize=11),
        ))

    doc.build(
        story,
        onFirstPage=lambda c, d: _cabecalho_rodape(c, d, "Grade Horária"),
        onLaterPages=lambda c, d: _cabecalho_rodape(c, d, "Grade Horária"),
    )
    buf.seek(0)
    return buf.getvalue()


def gerar_excel_grade(engine, ano_id, turma_ids=None):
    """Gera Excel com 1 aba consolidada + 1 aba por turma."""
    if turma_ids is None:
        turma_ids = _todas_turmas(engine, ano_id)

    wb = Workbook()
    wb.remove(wb.active)

    # ---------- Aba consolidada ----------
    ws_all = wb.create_sheet("Todas as aulas")
    ws_all.append(["Turma", "Série", "Dia", "Ordem",
                   "Início", "Fim", "Disciplina", "Professor"])
    for cell in ws_all[1]:
        cell.fill = XL_AZUL
        cell.font = XL_FONTE_HEADER
        cell.alignment = XL_CENTRO
        cell.border = XL_BORDA

    with engine.connect() as conn:
        todos = conn.execute(text("""
            SELECT t.codigo, t.serie, h.dia, h.ordem, h.inicio, h.fim,
                   c.nome AS componente, p.nome AS professor
            FROM grade_gerada g
            JOIN atividades a ON a.id = g.atividade_id
            JOIN turmas t ON t.id = a.turma_id
            JOIN horarios h ON h.id = g.horario_id
            JOIN componentes c ON c.id = a.componente_id
            JOIN professores p ON p.id = a.professor_id
            WHERE g.ano_letivo_id = :a
            ORDER BY t.codigo, h.dia, h.ordem
        """), {"a": ano_id}).fetchall()

    for row in todos:
        ws_all.append([row.codigo, row.serie, row.dia, row.ordem,
                      row.inicio, row.fim, row.componente, row.professor])

    larguras = [10, 10, 10, 8, 10, 10, 25, 25]
    for i, w in enumerate(larguras, start=1):
        ws_all.column_dimensions[get_column_letter(i)].width = w

    # ---------- Aba por turma ----------
    for tid in turma_ids:
        dados = _buscar_grade_turma(engine, ano_id, tid)
        if not dados:
            continue

        t = dados["turma"]
        slots = dados["slots"]
        mapa = dados["mapa"]

        nome_aba = f"{t.codigo}"[:31]
        ws = wb.create_sheet(nome_aba)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        c = ws.cell(row=1, column=1, value=f"Grade Horária — {t.nome}")
        c.font = Font(bold=True, size=14, color="1F4E78")
        c.alignment = XL_CENTRO

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
        c = ws.cell(row=2, column=1,
                    value=f"Série: {t.serie}  |  Turno: {t.turno}  |  Código: {t.codigo}")
        c.font = Font(italic=True, size=10, color="7F8C8D")
        c.alignment = XL_CENTRO

        ws.cell(row=4, column=1, value="Horário")
        for j, dia in enumerate(DIAS, start=2):
            ws.cell(row=4, column=j, value=dia.upper())

        for j in range(1, 7):
            cell = ws.cell(row=4, column=j)
            cell.fill = XL_AZUL
            cell.font = XL_FONTE_HEADER
            cell.alignment = XL_CENTRO
            cell.border = XL_BORDA

        for i, slot in enumerate(slots, start=5):
            c = ws.cell(row=i, column=1,
                       value=f"{slot.nome}\n{slot.inicio}-{slot.fim}")
            c.fill = XL_AZUL_CLARO
            c.font = XL_FONTE_BOLD
            c.alignment = XL_CENTRO
            c.border = XL_BORDA

            for j, dia in enumerate(DIAS, start=2):
                celula = mapa.get((dia, slot.ordem))
                cell = ws.cell(row=i, column=j)
                if celula:
                    comp, prof = celula
                    cell.value = f"{comp}\n({prof})"
                else:
                    cell.value = "—"
                cell.alignment = XL_CENTRO
                cell.border = XL_BORDA
                cell.font = XL_FONTE_NORMAL

            ws.row_dimensions[i].height = 40

        ws.column_dimensions["A"].width = 16
        for j in range(2, 7):
            ws.column_dimensions[get_column_letter(j)].width = 22

    if len(wb.sheetnames) == 0:
        wb.create_sheet("Vazio")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
